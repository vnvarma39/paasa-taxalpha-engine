import io
import csv
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ReportExporter:
    """
    Generates CA-Ready Excel Workbooks (.xlsx) and CSV reports for Indian Chartered Accountants.
    Includes Schedule Capital Gains, Schedule FSI (Dividends/Form 67), Schedule FA (Foreign Assets),
    and March 31 Tax-Loss Harvesting Plan.
    """

    @staticmethod
    def export_excel(analysis_data, selected_fy=None):
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style definitions
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Dark Slate
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sub_header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid") # Emerald
        sub_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        accent_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        bold_font = Font(name="Calibri", size=11, bold=True)
        thin_border = Border(
            left=Side(style='thin', color="E2E8F0"),
            right=Side(style='thin', color="E2E8F0"),
            top=Side(style='thin', color="E2E8F0"),
            bottom=Side(style='thin', color="E2E8F0")
        )

        realized = analysis_data.get("realized_trades", [])
        if selected_fy and selected_fy != "ALL":
            realized = [t for t in realized if t.get("indian_fy") == selected_fy]

        dividends = analysis_data.get("dividends", [])
        if selected_fy and selected_fy != "ALL":
            dividends = [d for d in dividends if d.get("indian_fy") == selected_fy]

        open_positions = analysis_data.get("open_positions", [])
        harvest = analysis_data.get("tax_loss_harvest", {})
        candidates = harvest.get("candidates", [])

        # ==========================================
        # TAB 1: SUMMARY DASHBOARD
        # ==========================================
        ws1 = wb.create_sheet(title="Summary_Dashboard")
        ws1.views.sheetView[0].showGridLines = True

        ws1["A1"] = "PAASA TAXALPHA ENGINE — CA TAX SUMMARY REPORT"
        ws1["A1"].font = title_font
        ws1["A2"] = f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')} | Applicable FY: {selected_fy or 'All Financial Years'}"
        ws1["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

        summary_rows = [
            ("Portfolio Current Market Value (USD)", f"${analysis_data.get('portfolio_totals', {}).get('current_val_usd', 0):,.2f}"),
            ("Portfolio Current Market Value (INR)", f"₹{analysis_data.get('portfolio_totals', {}).get('current_val_inr', 0):,.2f}"),
            ("Applicable Rule 115 Exchange Rate (Current)", f"₹{analysis_data.get('portfolio_totals', {}).get('current_fx_rate', 0):.2f} / USD"),
            ("", ""),
            ("Realized STCG (Short-Term Capital Gains INR)", f"₹{sum(r['gain_inr'] for r in realized if r['term']=='STCG'):,.2f}"),
            ("Realized LTCG (Long-Term Capital Gains INR)", f"₹{sum(r['gain_inr'] for r in realized if r['term']=='LTCG'):,.2f}"),
            ("Total Realized Capital Gains (INR)", f"₹{sum(r['gain_inr'] for r in realized):,.2f}"),
            ("Estimated Capital Gains Tax Liability (INR)", f"₹{sum(r['estimated_tax_inr'] for r in realized):,.2f}"),
            ("", ""),
            ("Gross Foreign Dividends Received (USD)", f"${sum(d['gross_div_usd'] for d in dividends):,.2f}"),
            ("Gross Foreign Dividends in INR (Rule 115)", f"₹{sum(d['gross_div_inr'] for d in dividends):,.2f}"),
            ("US Withholding Tax (25% DTAA Rate INR)", f"₹{sum(d['withholding_inr'] for d in dividends):,.2f}"),
            ("Foreign Tax Credit (FTC) Claimed under Sec 90", f"₹{sum(d.get('ftc_eligible_inr', 0) for d in dividends):,.2f}"),
            ("", ""),
            ("Pre-March 31 Potential Harvestable Loss (INR)", f"₹{harvest.get('total_harvestable_loss_inr', 0):,.2f}"),
            ("PRE-MARCH 31 TAX ALPHA SAVED (INR)", f"₹{harvest.get('total_potential_tax_saved_inr', 0):,.2f}")
        ]

        ws1["A4"] = "Key Financial Metric"
        ws1["B4"] = "Computed Value"
        ws1["A4"].fill = header_fill
        ws1["A4"].font = header_font
        ws1["B4"].fill = header_fill
        ws1["B4"].font = header_font

        row_idx = 5
        for item, val in summary_rows:
            ws1[f"A{row_idx}"] = item
            ws1[f"B{row_idx}"] = val
            if "TAX ALPHA SAVED" in item:
                ws1[f"A{row_idx}"].fill = sub_header_fill
                ws1[f"A{row_idx}"].font = sub_header_font
                ws1[f"B{row_idx}"].fill = sub_header_fill
                ws1[f"B{row_idx}"].font = sub_header_font
            elif item == "":
                pass
            else:
                ws1[f"A{row_idx}"].border = thin_border
                ws1[f"B{row_idx}"].border = thin_border
                ws1[f"B{row_idx}"].alignment = Alignment(horizontal="right")
            row_idx += 1

        # ==========================================
        # TAB 2: SCHEDULE CAPITAL GAINS (FIFO)
        # ==========================================
        ws2 = wb.create_sheet(title="Schedule_Capital_Gains")
        ws2.views.sheetView[0].showGridLines = True

        headers2 = [
            "Indian FY", "US CY", "Symbol", "Asset Type", "Qty", "Buy Date", "Sell Date",
            "Holding (Days)", "Holding (Mos)", "Term", "Buy Price ($)", "Sell Price ($)",
            "Buy FX (Rule 115)", "Sell FX (Rule 115)", "Cost Basis ($)", "Proceeds ($)",
            "Gain ($)", "Cost Basis (₹)", "Proceeds (₹)", "Gain (₹)", "Tax Rate (%)", "Estimated Tax (₹)"
        ]

        for col_idx, h in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, r in enumerate(realized, 2):
            vals = [
                r.get("indian_fy"), r.get("us_cy"), r.get("symbol"), r.get("asset_type"),
                r.get("quantity"), r.get("buy_date"), r.get("sell_date"),
                r.get("holding_days"), r.get("holding_months"), r.get("term"),
                r.get("buy_price_usd"), r.get("sell_price_usd"),
                r.get("buy_fx_rate"), r.get("sell_fx_rate"),
                r.get("buy_cost_usd"), r.get("sell_proceeds_usd"), r.get("gain_usd"),
                r.get("buy_cost_inr"), r.get("sell_proceeds_inr"), r.get("gain_inr"),
                r.get("tax_rate_pct"), r.get("estimated_tax_inr")
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws2.cell(row=r_idx, column=col_idx, value=val)
                c.border = thin_border
                if isinstance(val, (int, float)):
                    c.alignment = Alignment(horizontal="right")

        # ==========================================
        # TAB 3: SCHEDULE FSI (DIVIDENDS & FORM 67)
        # ==========================================
        ws3 = wb.create_sheet(title="Schedule_FSI_Dividends")
        ws3.views.sheetView[0].showGridLines = True

        headers3 = [
            "Date", "Symbol", "Indian FY", "Quantity", "Div/Share ($)",
            "Gross Div ($)", "US WHT 25% ($)", "Net Div ($)",
            "Rule 115 FX", "Gross Div (₹)", "US WHT (₹)", "Indian Tax (₹)",
            "FTC Claimed Sec 90 (₹)", "Net Tax Due (₹)", "Country Code"
        ]

        for col_idx, h in enumerate(headers3, 1):
            cell = ws3.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, d in enumerate(dividends, 2):
            vals = [
                d.get("date"), d.get("symbol"), d.get("indian_fy"), d.get("quantity"), d.get("per_share_usd"),
                d.get("gross_div_usd"), d.get("withholding_usd"), d.get("net_div_usd"),
                d.get("fx_rate"), d.get("gross_div_inr"), d.get("withholding_inr"), d.get("indian_tax_inr"),
                d.get("ftc_eligible_inr"), d.get("net_tax_payable_inr"), "01 (USA)"
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws3.cell(row=r_idx, column=col_idx, value=val)
                c.border = thin_border
                if isinstance(val, (int, float)):
                    c.alignment = Alignment(horizontal="right")

        # ==========================================
        # TAB 4: SCHEDULE FA (FOREIGN ASSETS FOR ITR-2)
        # ==========================================
        ws4 = wb.create_sheet(title="Schedule_FA_Holdings")
        ws4.views.sheetView[0].showGridLines = True

        headers4 = [
            "Lot ID", "Symbol", "Asset Type", "Buy Date", "Quantity",
            "Acquisition Cost ($)", "Acquisition Cost (₹)", "Buy FX (Rule 115)",
            "Current Price ($)", "Closing Value ($)", "Closing Value (₹)",
            "Unrealized P&L ($)", "Unrealized P&L (₹)", "Country Code", "Nature of Entity"
        ]

        for col_idx, h in enumerate(headers4, 1):
            cell = ws4.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, p in enumerate(open_positions, 2):
            vals = [
                p.get("lot_id"), p.get("symbol"), p.get("asset_type"), p.get("buy_date"), p.get("quantity"),
                p.get("cost_basis_usd"), p.get("cost_basis_inr"), p.get("buy_fx_rate"),
                p.get("current_price_usd"), p.get("current_val_usd"), p.get("current_val_inr"),
                p.get("unrealized_pnl_usd"), p.get("unrealized_pnl_inr"), "01 (USA)", "Foreign Listed Body Corporate"
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws4.cell(row=r_idx, column=col_idx, value=val)
                c.border = thin_border
                if isinstance(val, (int, float)):
                    c.alignment = Alignment(horizontal="right")

        # ==========================================
        # TAB 5: TAX ALPHA HARVEST PLAN
        # ==========================================
        ws5 = wb.create_sheet(title="Tax_Alpha_Harvest_Plan")
        ws5.views.sheetView[0].showGridLines = True

        headers5 = [
            "Symbol", "Quantity", "Current Price ($)", "Loss ($)", "Loss (₹)",
            "Term", "Tax Rate (%)", "Tax Saved (₹)",
            "Recommended Swap ETF", "Substitute Name", "Correlation", "Swap Action Note"
        ]

        for col_idx, h in enumerate(headers5, 1):
            cell = ws5.cell(row=1, column=col_idx, value=h)
            cell.fill = sub_header_fill
            cell.font = sub_header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, c_item in enumerate(candidates, 2):
            swap = c_item.get("swap", {})
            vals = [
                c_item.get("symbol"), c_item.get("quantity"), c_item.get("current_price_usd"),
                c_item.get("loss_usd"), c_item.get("loss_inr"),
                c_item.get("term"), c_item.get("applicable_tax_rate"), c_item.get("potential_tax_saved_inr"),
                swap.get("primary_substitute"), swap.get("substitute_name"), swap.get("correlation"),
                c_item.get("action_recommendation")
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws5.cell(row=r_idx, column=col_idx, value=val)
                c.border = thin_border
                if isinstance(val, (int, float)):
                    c.alignment = Alignment(horizontal="right")

        # Adjust column widths for all sheets
        for ws in [ws1, ws2, ws3, ws4, ws5]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_csv(analysis_data, selected_fy=None):
        """
        Exports a flat consolidated CSV of realized capital gains and Rule 115 conversions.
        """
        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow(["=== PAASA TAXALPHA ENGINE: SCHEDULE CAPITAL GAINS & RULE 115 CONVERSION ==="])
        writer.writerow([])
        writer.writerow([
            "Indian_FY", "US_CY", "Symbol", "Asset_Type", "Quantity", "Buy_Date", "Sell_Date",
            "Holding_Days", "Holding_Months", "Term", "Buy_Price_USD", "Sell_Price_USD",
            "Buy_FX_Rule115", "Sell_FX_Rule115", "Cost_Basis_USD", "Proceeds_USD",
            "Gain_USD", "Cost_Basis_INR", "Proceeds_INR", "Gain_INR", "Tax_Rate_Pct", "Estimated_Tax_INR"
        ])

        realized = analysis_data.get("realized_trades", [])
        if selected_fy and selected_fy != "ALL":
            realized = [t for t in realized if t.get("indian_fy") == selected_fy]

        for r in realized:
            writer.writerow([
                r.get("indian_fy"), r.get("us_cy"), r.get("symbol"), r.get("asset_type"),
                r.get("quantity"), r.get("buy_date"), r.get("sell_date"),
                r.get("holding_days"), r.get("holding_months"), r.get("term"),
                r.get("buy_price_usd"), r.get("sell_price_usd"),
                r.get("buy_fx_rate"), r.get("sell_fx_rate"),
                r.get("buy_cost_usd"), r.get("sell_proceeds_usd"), r.get("gain_usd"),
                r.get("buy_cost_inr"), r.get("sell_proceeds_inr"), r.get("gain_inr"),
                r.get("tax_rate_pct"), r.get("estimated_tax_inr")
            ])

        output.seek(0)
        return output.getvalue()
